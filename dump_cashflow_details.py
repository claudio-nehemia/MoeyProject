import openpyxl

wb = openpyxl.load_workbook('c:/projectFlutter/MOEYPROJECT/Simulasi Bisnis & Cashflow Moey.xlsx', data_only=False)
wb_val = openpyxl.load_workbook('c:/projectFlutter/MOEYPROJECT/Simulasi Bisnis & Cashflow Moey.xlsx', data_only=True)

ws = wb['Cashflow Contoh "Cust-Ambie"']
ws_val = wb_val['Cashflow Contoh "Cust-Ambie"']

with open('c:/projectFlutter/MOEYPROJECT/MoeyBackendAdmin/scratch_cashflow_details.txt', 'w', encoding='utf-8') as f:
    for r in range(1, 80):
        row_vals = []
        for c in range(28, min(ws.max_column + 1, 60)): # from col AB (28) onwards
            cell_formula = ws.cell(row=r, column=c).value
            cell_val = ws_val.cell(row=r, column=c).value
            if cell_formula is not None or cell_val is not None:
                col_letter = openpyxl.utils.get_column_letter(c)
                if str(cell_formula).startswith('='):
                    row_vals.append(f"{col_letter}{r}: [{cell_val}] ({cell_formula})")
                else:
                    row_vals.append(f"{col_letter}{r}: [{cell_val}]")
        if row_vals:
            f.write(f"Row {r:2d}: " + " | ".join(row_vals) + "\n")

print("DUMP COMPLETED")
