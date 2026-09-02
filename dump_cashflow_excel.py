import openpyxl
import json

wb = openpyxl.load_workbook('c:/projectFlutter/MOEYPROJECT/Simulasi Bisnis & Cashflow Moey.xlsx', data_only=False)
wb_val = openpyxl.load_workbook('c:/projectFlutter/MOEYPROJECT/Simulasi Bisnis & Cashflow Moey.xlsx', data_only=True)

with open('c:/projectFlutter/MOEYPROJECT/MoeyBackendAdmin/scratch_cashflow_sheets.txt', 'w', encoding='utf-8') as f:
    f.write("=== SHEET NAMES ===\n")
    for i, name in enumerate(wb.sheetnames):
        f.write(f"[{i}] {name}\n")
    
    for name in wb.sheetnames:
        f.write(f"\n\n==================== SHEET: {name} ====================\n")
        ws = wb[name]
        ws_val = wb_val[name]
        for r in range(1, min(ws.max_row + 1, 80)):
            row_vals = []
            has_val = False
            for c in range(1, min(ws.max_column + 1, 30)):
                cell_formula = ws.cell(row=r, column=c).value
                cell_val = ws_val.cell(row=r, column=c).value
                if cell_formula is not None or cell_val is not None:
                    has_val = True
                    col_letter = openpyxl.utils.get_column_letter(c)
                    if str(cell_formula).startswith('='):
                        row_vals.append(f"{col_letter}{r}: [{cell_val}] (FORMULA: {cell_formula})")
                    else:
                        row_vals.append(f"{col_letter}{r}: [{cell_val}]")
            if has_val:
                f.write(" | ".join(row_vals) + "\n")

print("DUMP COMPLETED")
